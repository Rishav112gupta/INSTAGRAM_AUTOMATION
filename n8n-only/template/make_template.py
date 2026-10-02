"""Creates instagram-posts-template.xlsx (upload it to Google Drive and open with Google Sheets).

    pip install openpyxl && python3 make_template.py
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

COLUMNS = [
    ("Post ID", 10, "Filled in automatically"),
    ("Topic", 34, "YOU: what the post is about"),
    ("Category", 16, "YOU (optional)"),
    ("Target Audience", 22, "YOU (optional)"),
    ("Important Info", 38, "YOU: real facts to include (dates, fees...)"),
    ("Call To Action", 26, "YOU (optional)"),
    ("Scheduled Date", 15, "Auto, or YOU: YYYY-MM-DD"),
    ("Post Time", 10, "Auto, or YOU: HH:MM"),
    ("Status", 16, "See Status list"),
    ("Headline", 30, "AI"),
    ("Caption", 60, "AI - you may edit"),
    ("Hashtags", 34, "AI - you may edit"),
    ("Image URL", 40, "AI (or paste your own JPEG link)"),
    ("Review Notes", 50, "Things to check before approving"),
    ("Instagram URL", 34, "Filled after publishing"),
    ("Instagram Media ID", 20, "Filled after publishing"),
    ("Container ID", 18, "Technical"),
    ("Error", 40, "Why something failed"),
    ("Last Updated", 17, "Automatic"),
]

SETTINGS = [
    ("Company Name", "Your Company Name", "Used in every post"),
    ("Company Description", "What the company does, in 1-3 sentences.", ""),
    ("Target Audience", "e.g. Class 11-12 students and their parents", "Default audience"),
    ("Brand Voice", "e.g. Friendly, encouraging, simple English", ""),
    ("Words To Avoid", "e.g. guaranteed, cheapest", ""),
    ("Default Hashtags", "yourbrand", "Added to every post (no #, separate with spaces)"),
    ("Content Rules", "e.g. Never mention competitors. No political content.", ""),
    ("Visual Style", "Clean flat illustrations, navy blue and orange colours", "Used for AI images"),
    ("Verified Facts", "Course fee (JEE crash course): Rs 4,999\nOffice: 12 Main Road, Pune\nPhone: +91 90000 00000", "The ONLY prices, dates and claims the AI may use. One per line."),
    ("Website", "https://www.example.com", ""),
    ("Contact Info", "", ""),
    ("Language", "English", ""),
    ("Post Every N Days", "4", "4 = every 4th day"),
    ("Start Date", "2026-10-05", "First posting date (YYYY-MM-DD). Dates follow this pattern."),
    ("Default Post Time", "18:00", "24-hour time"),
    ("Generate Days Before", "3", "How many days before the date the AI writes the post"),
    ("Timezone", "Asia/Kolkata", ""),
    ("Headline On Image", "No", "Yes = write the headline on the image (Cloudinary)"),
]

STATUS_HELP = [
    ("(empty) or NEW", "You added a topic. n8n will give it an ID and a date."),
    ("PLANNED", "Date assigned. The AI will write it a few days before."),
    ("NEEDS_REVIEW", "AI wrote it. READ IT, check Review Notes, edit if needed."),
    ("APPROVED", "YOU set this. It will be published at its date and time."),
    ("REGENERATE", "YOU set this to get a new AI version (write what to change in Review Notes)."),
    ("REJECTED", "YOU set this to drop the post. Nothing happens."),
    ("PUBLISHING", "n8n is posting it right now."),
    ("PUBLISHED", "Posted. Instagram URL is filled in."),
    ("FAILED", "Posting failed - see Error. Fix, then set APPROVED again to retry."),
    ("PROBLEM", "Approved but something is missing - see Error."),
]

HEADER = PatternFill("solid", fgColor="1E3A8A")
WHITE_BOLD = Font(bold=True, color="FFFFFF")


def build(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Posts"
    for i, (name, width, _) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=i, value=name)
        c.fill, c.font = HEADER, WHITE_BOLD
        ws.column_dimensions[c.column_letter].width = width
    ws.freeze_panes = "C2"
    # Keep dates/times as plain text so nothing gets converted by locale.
    for col in ("A", "G", "H", "P", "Q"):
        for r in range(2, 1001):
            ws[f"{col}{r}"].number_format = "@"
    examples = [
        {"Topic": "Last-minute exam preparation tips", "Category": "EXAM", "Target Audience": "Class 12 students", "Call To Action": "Follow us for daily tips"},
        {"Topic": "Welcome message for the new weekend batch", "Category": "BATCH", "Important Info": "Weekend batch starts Saturday 10 AM"},
        {"Topic": "Myth vs fact: studying long hours", "Category": "NORMAL"},
    ]
    for r, ex in enumerate(examples, start=2):
        for i, (name, _, _) in enumerate(COLUMNS, start=1):
            if name in ex:
                ws.cell(row=r, column=i, value=ex[name])
    dv = DataValidation(type="list", formula1='"NEW,PLANNED,NEEDS_REVIEW,APPROVED,REGENERATE,REJECTED,PUBLISHING,PUBLISHED,FAILED,PROBLEM"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("I2:I1000")
    for r in range(2, 1001):
        for col in ("E", "K", "N", "R"):
            ws[f"{col}{r}"].alignment = Alignment(wrap_text=True, vertical="top")

    st = wb.create_sheet("Settings")
    for i, h in enumerate(("Setting", "Value", "Help"), start=1):
        c = st.cell(row=1, column=i, value=h)
        c.fill, c.font = HEADER, WHITE_BOLD
    st.column_dimensions["A"].width = 24
    st.column_dimensions["B"].width = 60
    st.column_dimensions["C"].width = 60
    for r, (k, v, h) in enumerate(SETTINGS, start=2):
        st.cell(row=r, column=1, value=k).font = Font(bold=True)
        cell = st.cell(row=r, column=2, value=v)
        cell.number_format = "@"
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        st.cell(row=r, column=3, value=h)

    hs = wb.create_sheet("How to use")
    hs.column_dimensions["A"].width = 18
    hs.column_dimensions["B"].width = 90
    hs.cell(row=1, column=1, value="Status").font = Font(bold=True)
    hs.cell(row=1, column=2, value="Meaning").font = Font(bold=True)
    for r, (k, v) in enumerate(STATUS_HELP, start=2):
        hs.cell(row=r, column=1, value=k).font = Font(bold=True)
        hs.cell(row=r, column=2, value=v)
    tips = [
        "",
        "Daily routine:",
        "1. Add topics as new rows in the Posts tab (only Topic is required).",
        "2. Filter Status = NEEDS_REVIEW. Read caption + Review Notes, open the Image URL, edit if needed.",
        "3. Set Status to APPROVED (or REGENERATE / REJECTED).",
        "4. Approved posts are published automatically at their Scheduled Date + Post Time.",
        "Do not rename the column headers or the tab names (Posts, Settings).",
    ]
    for i, t in enumerate(tips, start=len(STATUS_HELP) + 2):
        hs.cell(row=i, column=1, value=t)
    wb.save(path)


if __name__ == "__main__":
    out = Path(__file__).with_name("instagram-posts-template.xlsx")
    build(out)
    print("wrote", out)
